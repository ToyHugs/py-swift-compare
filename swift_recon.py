#!/usr/bin/env python3
"""
Réconciliation de messages SWIFT MT535 entre deux bases :
CTRL = ancienne version, TEST = nouvelle version.

Étapes :
  1. Extraction des messages de chaque base, filtrés sur une plage de dates
     (colonne timestamp) propre à chaque base.
  2. Parsing : chaque champ de chaque message devient une ligne
     (compte, date de relevé, ISIN, bloc, tag, qualifiant, valeur)
     stockée dans une base SQLite locale.
  3. Comparaison champ à champ, en SQL dans SQLite.
  4. Export Excel : synthèse, comptes, positions, écarts détaillés, erreurs.

Usage :
  python swift_recon.py config.ini
  python swift_recon.py config.ini --from 2026-10-01 --to 2026-10-21
  python swift_recon.py config.ini --ctrl-from 2026-10-20 --ctrl-to 2026-10-20 \\
                                   --test-from 2026-10-22 --test-to 2026-10-22
  python swift_recon.py config.ini --compare-only   # rejoue la comparaison sans ré-extraire
"""
from __future__ import annotations

import argparse
import configparser
import csv
import getpass
import os
import re
import sqlite3
import sys
import time
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Iterator

SIDES = ("CTRL", "TEST")
EXCEL_MAX_ROWS = 1_048_575
DATE_FMT = "%Y-%m-%d %H:%M:%S"


# =============================================================================
# 1. Parsing SWIFT
# =============================================================================

FIELD_RE = re.compile(r"^:(\d{2}[A-Z]?):(.*)$")           # :93B:...
GENERIC_RE = re.compile(r"^:([A-Z0-9]{4})/(.*)$", re.S)    # :AGGR//UNIT/100,
BLOCK2_TYPE_RE = re.compile(r"\{2:[IO](\d{3})")
ISIN_RE = re.compile(r"ISIN ([A-Z0-9]{12})")
OTHER_ID_RE = re.compile(r"^(/[A-Z0-9]{2}/.+)$", re.M)     # /XX/code-interne
SPACES_RE = re.compile(r"\s+")
TRAILING_ZEROS_RE = re.compile(r"(,\d*?)0+$")


@dataclass
class Field:
    path: str   # blocs 16R imbriqués, ex. "SUBSAFE/FIN/SUBBAL"
    tag: str    # ex. "93B"
    qual: str   # ex. "AGGR" ("" si le champ n'a pas de qualifiant)
    value: str  # valeur brute, sans le qualifiant


def message_type(text: str) -> str | None:
    """Type de message lu dans le bloc 2 ({2:I535...}), None si absent."""
    m = BLOCK2_TYPE_RE.search(text[:300])
    return m.group(1) if m else None


def raw_fields(text: str) -> Iterator[tuple[str, str]]:
    """(tag, valeur) de chaque champ du bloc 4. Les valeurs multi-lignes sont jointes par \\n."""
    start = text.find("{4:")
    body = text[start + 3:] if start >= 0 else text
    tag, lines = None, []
    for line in body.splitlines():
        line = line.rstrip()
        if line.startswith("-}"):
            break
        m = FIELD_RE.match(line)
        if m:
            if tag:
                yield tag, "\n".join(lines)
            tag, lines = m.group(1), [m.group(2)]
        elif tag and line:
            lines.append(line)
    if tag:
        yield tag, "\n".join(lines)


def parse_fields(text: str) -> list[Field]:
    """Champs du message avec leur bloc. 16R/16S servent uniquement à suivre l'imbrication."""
    path: list[str] = []
    fields = []
    for tag, value in raw_fields(text):
        if tag == "16R":
            path.append(value.strip())
        elif tag == "16S":
            if path:
                path.pop()
        else:
            qual = ""
            m = GENERIC_RE.match(value)
            if m:
                qual, value = m.group(1), m.group(2).lstrip("/")
            fields.append(Field("/".join(path), tag, qual, value))
    return fields


def normalize(value: str) -> str:
    """Neutralise les différences de forme : espaces et retours ligne,
    '1000,00' / '1000,' -> '1000', '12,50' -> '12,5'."""
    v = SPACES_RE.sub(" ", value).strip()
    v = TRAILING_ZEROS_RE.sub(r"\1", v)
    return v[:-1] if v.endswith(",") else v


def security_id(value_35b: str) -> str | None:
    """ISIN si présent, sinon identifiant /XX/... (code propriétaire ou national)."""
    m = ISIN_RE.search(value_35b)
    if m:
        return m.group(1)
    m = OTHER_ID_RE.search(value_35b)
    return m.group(1).strip() if m else None


@dataclass
class Statement:
    account: str | None
    stmt_date: str | None
    page: str | None
    stmt_no: str | None
    rows: list[tuple]   # (isin, bloc, tag, qualifiant, valeur normalisée)


def explode_535(text: str) -> Statement:
    """Découpe un MT535 : entête (bloc GENL) + une ligne par champ,
    chaque champ d'un bloc FIN portant l'ISIN de sa position."""
    fields = parse_fields(text)

    head: dict[str, str] = {}
    for f in fields:
        if f.path != "GENL":
            continue
        if f.tag[:2] == "97" and f.qual == "SAFE":
            head.setdefault("account", f.value.split("\n")[0].strip())
        elif f.tag[:2] == "98" and f.qual == "STAT":
            head.setdefault("stmt_date", f.value[:8])
        elif f.tag == "28E":
            head["page"] = f.value.strip()
        elif f.tag == "13A" and f.qual == "STAT":
            head["stmt_no"] = f.value.strip()

    rows, isin = [], None
    for f in fields:
        if "FIN" not in f.path.split("/"):
            isin = None
        elif f.tag == "35B":          # toujours le premier champ d'un bloc FIN
            isin = security_id(f.value)
        rows.append((isin, f.path, f.tag, f.qual, normalize(f.value)))

    return Statement(head.get("account"), head.get("stmt_date"),
                     head.get("page"), head.get("stmt_no"), rows)


# =============================================================================
# 2. Sources de messages
# =============================================================================

def parse_date(text: str, end: bool) -> datetime:
    """'2026-10-21', '2026-10-21 14:30' ou '2026-10-21 14:30:00'.
    Pour une date de fin, retourne la borne exclusive (fin incluse dans la plage)."""
    text = text.strip()
    for fmt, step in (("%Y-%m-%d %H:%M:%S", timedelta(seconds=1)),
                      ("%Y-%m-%d %H:%M", timedelta(minutes=1)),
                      ("%Y-%m-%d", timedelta(days=1))):
        try:
            d = datetime.strptime(text, fmt)
        except ValueError:
            continue
        return d + step if end else d
    raise ValueError(f"Date invalide : {text!r} (attendu AAAA-MM-JJ[ HH:MM[:SS]])")


def get_password(sec: configparser.SectionProxy, side: str) -> str:
    """Variable d'environnement SWIFT_CTRL_PASSWORD / SWIFT_TEST_PASSWORD, sinon saisie."""
    return (os.environ.get(f"SWIFT_{side}_PASSWORD")
            or sec.get("password", "")
            or getpass.getpass(f"Mot de passe {side} ({sec['user']}@{sec['dsn']}) : "))


_thick_mode_on = False


def oracle_messages(sec: configparser.SectionProxy, password: str,
                    d_from: datetime, d_to: datetime, mtype: str) -> Iterator[tuple[str, str]]:
    """(msg_id, texte) des messages de la plage [d_from, d_to[ sur la colonne timestamp."""
    import oracledb

    global _thick_mode_on
    if sec.getboolean("thick_mode", fallback=False) and not _thick_mode_on:
        oracledb.init_oracle_client(lib_dir=sec.get("instant_client_dir") or None)
        _thick_mode_on = True
    oracledb.defaults.fetch_lobs = False     # CLOB récupérés directement en str

    ts = sec.get("ts_column") or '"TIMESTAMP"'
    sql = (f"SELECT {sec['id_column']}, {sec['msg_column']} FROM {sec['table']} "
           f"WHERE {ts} >= TO_DATE(:d1, 'YYYY-MM-DD HH24:MI:SS') "
           f"AND {ts} < TO_DATE(:d2, 'YYYY-MM-DD HH24:MI:SS')")
    binds = {"d1": d_from.strftime(DATE_FMT), "d2": d_to.strftime(DATE_FMT)}
    if sec.get("type_column"):
        sql += f" AND {sec['type_column']} LIKE :mtype"
        binds["mtype"] = f"%{mtype}"
    if sec.get("where"):
        sql += f" AND ({sec['where']})"

    with oracledb.connect(user=sec["user"], password=password, dsn=sec["dsn"]) as conn:
        with conn.cursor() as cur:
            cur.arraysize = 500
            cur.prefetchrows = 501
            cur.execute(sql, binds)
            for msg_id, text in cur:
                yield str(msg_id), text or ""


def file_messages(sec: configparser.SectionProxy, base_dir: Path) -> Iterator[tuple[str, str]]:
    """Mode fichiers (démo, messages exportés) : un message par fichier.
    Le dossier est relatif au fichier de configuration."""
    folder = base_dir / sec["dir"]
    if not folder.is_dir():
        sys.exit(f"Dossier de messages introuvable : {folder}")
    for p in sorted(folder.glob(sec.get("pattern", "*.txt"))):
        yield p.name, p.read_text(encoding="utf-8", errors="replace")


# =============================================================================
# 3. Stockage SQLite et comparaison
# =============================================================================

SCHEMA = """
PRAGMA journal_mode = OFF;
PRAGMA synchronous = OFF;
PRAGMA cache_size = -200000;
CREATE TABLE meta (k TEXT PRIMARY KEY, v TEXT);
CREATE TABLE msg  (side TEXT, msg_id TEXT, account TEXT, stmt_date TEXT, page TEXT,
                   stmt_no TEXT, nb_fields INTEGER, error TEXT);
CREATE TABLE fld  (side TEXT, msg_id TEXT, account TEXT, stmt_date TEXT, isin TEXT,
                   path TEXT, tag TEXT, qual TEXT, val TEXT);
"""


def extract(db: sqlite3.Connection, side: str,
            messages: Iterator[tuple[str, str]], mtype: str) -> tuple[int, int]:
    """Parse et stocke les messages d'une base. Retourne (nb chargés, nb ignorés car autre type)."""
    loaded = skipped = 0
    batch: list[tuple] = []
    t0 = time.time()
    for msg_id, text in messages:
        detected = message_type(text)
        if detected and detected != mtype:
            skipped += 1
            continue
        loaded += 1
        try:
            st = explode_535(text)
            err = None
            if not st.rows:
                err = "aucun champ trouvé"
            elif not st.account or not st.stmt_date:
                err = "compte (97A::SAFE) ou date de relevé (98a::STAT) introuvable"
        except Exception as exc:  # un message illisible ne doit pas arrêter le traitement
            st, err = Statement(None, None, None, None, []), f"{type(exc).__name__}: {exc}"

        db.execute("INSERT INTO msg VALUES (?,?,?,?,?,?,?,?)",
                   (side, msg_id, st.account, st.stmt_date, st.page, st.stmt_no, len(st.rows), err))
        batch.extend((side, msg_id, st.account, st.stmt_date, *r) for r in st.rows)
        if len(batch) >= 50_000:
            db.executemany("INSERT INTO fld VALUES (?,?,?,?,?,?,?,?,?)", batch)
            batch.clear()
        if loaded % 1000 == 0:
            print(f"  {side} : {loaded} messages ({time.time() - t0:.0f} s)")

    db.executemany("INSERT INTO fld VALUES (?,?,?,?,?,?,?,?,?)", batch)
    db.commit()
    print(f"  {side} : {loaded} messages chargés, {skipped} d'un autre type ignorés "
          f"({time.time() - t0:.0f} s)")
    return loaded, skipped


def ignore_clause(rules: list[str]) -> tuple[str, list[str]]:
    """'28E' exclut un tag, '20C:SEME' exclut un tag + qualifiant."""
    sql, params = [], []
    for rule in rules:
        tag, _, qual = rule.partition(":")
        if qual:
            sql.append("NOT (tag = ? AND qual = ?)")
            params += [tag, qual]
        else:
            sql.append("tag <> ?")
            params.append(tag)
    return (" AND ".join(sql) or "1 = 1"), params


KEY = ("account", "stmt_date", "isin", "path", "tag", "qual", "occ")


def compare(db: sqlite3.Connection, ignore: list[str]) -> None:
    """Construit la table 'ecart' : un enregistrement par champ différent, manquant ou en trop."""
    where, params = ignore_clause(ignore)
    db.executescript("DROP TABLE IF EXISTS f; DROP TABLE IF EXISTS ecart;")

    # Valeurs distinctes par clé : l'entête répété sur chaque page n'est compté qu'une fois.
    # occ numérote les valeurs quand un même tag/qualifiant revient plusieurs fois.
    db.execute(f"""
        CREATE TABLE f AS
        SELECT g.*, ROW_NUMBER() OVER (
                   PARTITION BY side, account, stmt_date, isin, path, tag, qual
                   ORDER BY val) AS occ
        FROM (SELECT side,
                     IFNULL(account, '?')  AS account,
                     IFNULL(stmt_date, '?') AS stmt_date,
                     IFNULL(isin, '-')     AS isin,
                     path, tag, qual, val,
                     MIN(msg_id)           AS msg_id
              FROM fld
              WHERE {where}
              GROUP BY side, IFNULL(account, '?'), IFNULL(stmt_date, '?'),
                       IFNULL(isin, '-'), path, tag, qual, val) g""", params)
    db.execute("CREATE INDEX f_key ON f (side, account, stmt_date, isin, path, tag, qual, occ)")

    on = " AND ".join(f"o.{k} = s.{k}" for k in KEY)
    db.execute(f"""
        CREATE TABLE ecart AS
        SELECT s.account, s.stmt_date, s.isin, s.path, s.tag, s.qual,
               s.val AS ctrl_val, o.val AS test_val,
               s.msg_id AS ctrl_msg, o.msg_id AS test_msg,
               CASE WHEN o.side IS NULL THEN 'MANQUANT DANS TEST' ELSE 'ECART' END AS statut
        FROM f s LEFT JOIN f o ON o.side = 'TEST' AND {on}
        WHERE s.side = 'CTRL' AND (o.side IS NULL OR o.val IS NOT s.val)
        UNION ALL
        SELECT s.account, s.stmt_date, s.isin, s.path, s.tag, s.qual,
               NULL, s.val, NULL, s.msg_id, 'EN TROP DANS TEST'
        FROM f s LEFT JOIN f o ON o.side = 'CTRL' AND {on}
        WHERE s.side = 'TEST' AND o.side IS NULL""")
    db.commit()


# =============================================================================
# 4. Rapport Excel
# =============================================================================

Q_POSITIONS = """
SELECT k.account, k.stmt_date, k.isin,
       CASE WHEN k.in_ctrl AND k.in_test THEN
                 CASE WHEN e.n IS NULL THEN 'OK' ELSE 'ECART' END
            WHEN k.in_ctrl THEN 'MANQUANT DANS TEST'
            ELSE 'EN TROP DANS TEST' END              AS statut,
       IFNULL(e.n, 0)                                 AS nb_champs,
       CASE WHEN k.in_ctrl AND k.in_test THEN e.tags END AS champs_en_ecart
FROM (SELECT account, stmt_date, isin,
             MAX(side = 'CTRL') AS in_ctrl, MAX(side = 'TEST') AS in_test
      FROM f WHERE isin <> '-' GROUP BY account, stmt_date, isin) k
LEFT JOIN (SELECT account, stmt_date, isin, COUNT(*) AS n,
                  GROUP_CONCAT(DISTINCT tag || CASE WHEN qual <> '' THEN '::' || qual ELSE '' END) AS tags
           FROM ecart WHERE isin <> '-' GROUP BY account, stmt_date, isin) e
  ON e.account = k.account AND e.stmt_date = k.stmt_date AND e.isin = k.isin
ORDER BY statut = 'OK', k.account, k.stmt_date, k.isin
"""

Q_ACCOUNTS = """
WITH m AS (SELECT IFNULL(account, '?') AS account, IFNULL(stmt_date, '?') AS stmt_date,
                  SUM(side = 'CTRL') AS msg_ctrl, SUM(side = 'TEST') AS msg_test
           FROM msg GROUP BY 1, 2),
     p AS (SELECT account, stmt_date,
                  COUNT(DISTINCT CASE WHEN side = 'CTRL' THEN isin END) AS pos_ctrl,
                  COUNT(DISTINCT CASE WHEN side = 'TEST' THEN isin END) AS pos_test
           FROM f WHERE isin <> '-' GROUP BY account, stmt_date),
     e AS (SELECT account, stmt_date,
                  SUM(isin = '-') AS ec_entete, SUM(isin <> '-') AS ec_positions
           FROM ecart GROUP BY account, stmt_date)
SELECT m.account, m.stmt_date,
       CASE WHEN m.msg_test = 0 THEN 'MANQUANT DANS TEST'
            WHEN m.msg_ctrl = 0 THEN 'EN TROP DANS TEST'
            WHEN IFNULL(e.ec_entete, 0) + IFNULL(e.ec_positions, 0) = 0 THEN 'OK'
            ELSE 'ECART' END AS statut,
       m.msg_ctrl, m.msg_test,
       IFNULL(p.pos_ctrl, 0), IFNULL(p.pos_test, 0),
       IFNULL(e.ec_entete, 0), IFNULL(e.ec_positions, 0)
FROM m LEFT JOIN p USING (account, stmt_date) LEFT JOIN e USING (account, stmt_date)
ORDER BY statut = 'OK', m.account, m.stmt_date
"""

Q_ECARTS = """
SELECT statut, account, stmt_date, isin, path, tag, qual, ctrl_val, test_val, ctrl_msg, test_msg
FROM ecart
ORDER BY account, stmt_date, isin = '-' DESC, isin, path, tag, qual
"""

Q_ERRORS = """
SELECT side, msg_id, account, stmt_date, page, error
FROM msg WHERE error IS NOT NULL ORDER BY side, msg_id
"""


def _clean(value):
    """Supprime les caractères de contrôle refusés par Excel."""
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    return ILLEGAL_CHARACTERS_RE.sub("", value) if isinstance(value, str) else value


def write_sheet(wb, title: str, header: list[str], rows, widths: list[int]) -> int:
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = wb.create_sheet(title)
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    cells = []
    for h in header:
        c = WriteOnlyCell(ws, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="305496")
        cells.append(c)
    ws.append(cells)
    n = 0
    for row in rows:
        if n >= EXCEL_MAX_ROWS:
            break
        ws.append([_clean(v) for v in row])
        n += 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{n + 1}"
    return n


def write_report(db: sqlite3.Connection, out: Path) -> None:
    from openpyxl import Workbook

    one = lambda sql: db.execute(sql).fetchone()[0]  # noqa: E731
    meta = dict(db.execute("SELECT k, v FROM meta"))
    nb_ecarts = one("SELECT COUNT(*) FROM ecart")

    summary = [("Exécution", datetime.now().strftime(DATE_FMT))]
    for side in SIDES:
        summary.append((f"Période {side}", meta.get(f"period_{side}", "-")))
    summary.append(("Champs exclus", meta.get("ignore", "")))
    summary.append(("", ""))
    for side in SIDES:
        summary.append((f"Messages {side}", one(f"SELECT COUNT(*) FROM msg WHERE side = '{side}'")))
        summary.append((f"Messages {side} en erreur",
                        one(f"SELECT COUNT(*) FROM msg WHERE side = '{side}' AND error IS NOT NULL")))
        summary.append((f"Positions {side}", one(
            f"SELECT COUNT(*) FROM (SELECT DISTINCT account, stmt_date, isin "
            f"FROM f WHERE side = '{side}' AND isin <> '-')")))
    summary.append(("", ""))
    for statut, n in db.execute(f"SELECT statut, COUNT(*) FROM ({Q_POSITIONS}) GROUP BY statut"):
        summary.append((f"Positions — {statut}", n))
    for statut, n in db.execute("SELECT statut, COUNT(*) FROM ecart GROUP BY statut"):
        summary.append((f"Champs — {statut}", n))
    summary.append(("Champs — écarts dans l'entête",
                    one("SELECT COUNT(*) FROM ecart WHERE isin = '-'")))

    if nb_ecarts > EXCEL_MAX_ROWS:
        csv_path = out.with_name(out.stem + "_ecarts.csv")
        with csv_path.open("w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh, delimiter=";")
            w.writerow(["statut", "compte", "date", "isin", "bloc", "tag", "qualifiant",
                        "valeur_ctrl", "valeur_test", "msg_ctrl", "msg_test"])
            w.writerows(db.execute(Q_ECARTS))
        summary.append(("ATTENTION", f"Trop d'écarts pour Excel : liste complète dans {csv_path.name}"))

    wb = Workbook(write_only=True)
    write_sheet(wb, "Synthese", ["Indicateur", "Valeur"], summary, [40, 60])
    write_sheet(wb, "Comptes",
                ["compte", "date", "statut", "msg_ctrl", "msg_test", "positions_ctrl",
                 "positions_test", "ecarts_entete", "ecarts_positions"],
                db.execute(Q_ACCOUNTS), [20, 11, 22, 10, 10, 15, 15, 14, 16])
    write_sheet(wb, "Positions",
                ["compte", "date", "isin", "statut", "nb_champs", "champs_en_ecart"],
                db.execute(Q_POSITIONS), [20, 11, 22, 22, 11, 50])
    write_sheet(wb, "Ecarts",
                ["statut", "compte", "date", "isin", "bloc", "tag", "qualifiant",
                 "valeur_ctrl", "valeur_test", "msg_ctrl", "msg_test"],
                db.execute(Q_ECARTS), [20, 18, 11, 22, 20, 6, 10, 40, 40, 20, 20])
    write_sheet(wb, "Erreurs", ["base", "msg_id", "compte", "date", "page", "erreur"],
                db.execute(Q_ERRORS), [8, 20, 20, 11, 10, 70])
    wb.save(out)


# =============================================================================
# 5. Programme principal
# =============================================================================

def main(argv: list[str] | None = None) -> None:
    ap = argparse.ArgumentParser(description="Réconciliation SWIFT MT535 CTRL / TEST")
    ap.add_argument("config", help="fichier de configuration (.ini)")
    ap.add_argument("--from", dest="d_from", help="date de début pour les deux bases")
    ap.add_argument("--to", dest="d_to", help="date de fin (incluse) pour les deux bases")
    for side in ("ctrl", "test"):
        ap.add_argument(f"--{side}-from", help=f"date de début pour {side.upper()}")
        ap.add_argument(f"--{side}-to", help=f"date de fin (incluse) pour {side.upper()}")
    ap.add_argument("--compare-only", action="store_true",
                    help="réutilise la base SQLite existante, sans ré-extraire")
    args = ap.parse_args(argv)

    if sqlite3.sqlite_version_info < (3, 25, 0):
        sys.exit(f"SQLite {sqlite3.sqlite_version} trop ancien (3.25 minimum) : mettez Python à jour.")

    cfg = configparser.ConfigParser(interpolation=None)
    if not cfg.read(args.config, encoding="utf-8"):
        sys.exit(f"Configuration introuvable : {args.config}")
    run = cfg["run"]
    mtype = run.get("msg_type", "535").strip()
    if mtype != "535":
        sys.exit("Seul le MT535 est géré pour l'instant.")
    base_dir = Path(args.config).resolve().parent
    store_path = base_dir / run.get("sqlite", "recon_mt535.sqlite")
    out_path = base_dir / run.get("output", "recon_mt535.xlsx")
    ignore = [r.strip() for r in run.get("ignore", "").split(",") if r.strip()]

    if args.compare_only:
        if not store_path.exists():
            sys.exit(f"Base {store_path} absente : lancez d'abord sans --compare-only.")
        db = sqlite3.connect(store_path)
    else:
        # Plages de dates et mots de passe demandés avant de lancer les extractions
        jobs = []
        for side in SIDES:
            sec = cfg[side.lower()]
            if sec.get("dir"):
                jobs.append((side, sec, None, None, f"fichiers {sec['dir']}"))
                continue
            low = side.lower()
            d1 = getattr(args, f"{low}_from") or args.d_from or sec.get("date_from") or run.get("date_from")
            d2 = getattr(args, f"{low}_to") or args.d_to or sec.get("date_to") or run.get("date_to")
            if not (d1 and d2):
                sys.exit(f"[{low}] : date de début et de fin obligatoires (config ou --{low}-from/--{low}-to)")
            period = (parse_date(d1, end=False), parse_date(d2, end=True))
            label = f"{period[0]:%Y-%m-%d %H:%M:%S} → {period[1]:%Y-%m-%d %H:%M:%S} (exclu)"
            jobs.append((side, sec, period, get_password(sec, side), label))

        if store_path.exists():
            store_path.unlink()
        db = sqlite3.connect(store_path)
        db.executescript(SCHEMA)
        for side, sec, period, password, label in jobs:
            print(f"Extraction {side} : {label}")
            db.execute("INSERT INTO meta VALUES (?, ?)", (f"period_{side}", label))
            src = file_messages(sec, base_dir) if period is None else oracle_messages(sec, password, *period, mtype)
            extract(db, side, src, mtype)

    db.execute("INSERT OR REPLACE INTO meta VALUES ('ignore', ?)", (", ".join(ignore),))
    print("Comparaison…")
    compare(db, ignore)
    print("Export Excel…")
    write_report(db, out_path)
    db.close()
    print(f"Terminé : {out_path}")


if __name__ == "__main__":
    main()
